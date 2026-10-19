"""Build the MLOps Assignment 1 report (.docx) from the project's own files.

Pulls real numbers from:   data/processed/heart_clean.csv, reports/model_comparison.csv,
                           models/model_metadata.json, tests/*.py
Pulls images from:         reports/figures/*.png, screenshots/*.png
Anything missing is shown as a yellow [TODO] / placeholder box, never silently skipped.

Usage (from project root):
    pip install python-docx
    python scripts/build_report.py --repo-url https://github.com/<user>/heart-disease-mlops \
        --student-id <ID> --svc-port 8081
"""
import argparse
import json
import re
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402

# ----------------------------------------------------------------------------------
# EDIT THESE if you changed them while making the project your own
# ----------------------------------------------------------------------------------
GRIDS = {
    "Logistic Regression": "C in {0.01, 0.1, 1, 10}; penalty = l2",
    "Random Forest": "n_estimators in {100, 200}; max_depth in {None, 4, 8}; "
                     "min_samples_leaf in {1, 3}",
}

# expected screenshots: key -> (file stem inside screenshots/, caption)
SHOTS = {
    "mlflow_runs": ("mlflow_runs", "MLflow experiment with parent runs and nested candidate runs"),
    "mlflow_compare": ("mlflow_compare", "MLflow run comparison of the two models"),
    "mlflow_artifacts": ("mlflow_artifacts", "MLflow artifacts: plots/ and model/"),
    "predict_cli": ("predict_cli", "Command-line inference with src/predict.py"),
    "clean_env": ("clean_env_run", "Full pipeline executed in a clean virtual environment"),
    "local_pytest": ("local_pytest", "Local pytest run"),
    "ci_green": ("ci_green_run", "GitHub Actions run: lint, test and train jobs all passing"),
    "ci_pytest": ("ci_pytest_logs", "Expanded pytest step logs in GitHub Actions"),
    "ci_artifacts": ("ci_artifacts", "Artifacts uploaded by a workflow run"),
    "swagger_local": ("swagger_predict", "Swagger UI calling POST /predict"),
    "k8s_get_all": ("k8s_get_all", "kubectl get all -n heart-api"),
    "k8s_describe": ("k8s_describe", "kubectl describe deployment"),
    "k8s_logs": ("request_logs", "Structured JSON request logs from both pods"),
    "grafana": ("grafana_dashboard", "Grafana dashboard under generated traffic"),
    "metrics_raw": ("metrics_raw", "Raw /metrics output"),
    "pods_all": ("pods_all", "All pods across the heart-api and monitoring namespaces"),
}

TODO_COUNT = 0
MISSING = []


# ----------------------------------------------------------------------------------
# docx helpers
# ----------------------------------------------------------------------------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_borders(cell, color="BBBBBB"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Page ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


class Report:
    def __init__(self):
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
        add_page_number(s)

        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(5)
        for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
            h = self.doc.styles[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            h.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        self.fig_n = 0
        self.tab_n = 0

    # text ----------------------------------------------------------------------
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text="", bold=False, italic=False, size=None):
        para = self.doc.add_paragraph()
        self._runs(para, text, bold, italic, size)
        return para

    def _runs(self, para, text, bold=False, italic=False, size=None):
        """Supports **bold** inline and [TODO: ...] highlighting."""
        global TODO_COUNT
        parts = re.split(r"(\[TODO:[^\]]*\]|\*\*[^*]+\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("[TODO:"):
                run = para.add_run(part)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                TODO_COUNT += 1
            elif part.startswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            else:
                run = para.add_run(part)
                run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)

    def bullets(self, items):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, it)

    def code(self, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        shade(c, "F2F2F2")
        set_cell_borders(c, "DDDDDD")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            para = c.paragraphs[0] if i == 0 else c.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths=None, caption=None, font=9):
        if caption:
            self.tab_n += 1
            cp = self.doc.add_paragraph()
            r = cp.add_run(f"Table {self.tab_n}. {caption}")
            r.bold = True
            r.font.size = Pt(9.5)
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(2)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            shade(c, "1F3A5F")
            set_cell_borders(c)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            run = c.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.size = Pt(font)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                set_cell_borders(cells[i])
                if ri % 2:
                    shade(cells[i], "F5F8FC")
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                self._runs(para, str(val), size=font)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # images --------------------------------------------------------------------
    def figure(self, path, caption, width_in=5.6):
        self.fig_n += 1
        path = Path(path)
        if path.exists():
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(path), width=Inches(width_in))
        else:
            self._placeholder(f"Image missing: {path.name}")
            MISSING.append(str(path))
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Figure {self.fig_n}. {caption}")
        r.italic = True
        r.font.size = Pt(9)

    def figure_pair(self, left, right, caption, width_in=3.1):
        """Two images side by side."""
        self.fig_n += 1
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, path in zip(t.rows[0].cells, (left, right)):
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if Path(path).exists():
                para.add_run().add_picture(str(path), width=Inches(width_in))
            else:
                run = para.add_run(f"[Image missing: {Path(path).name}]")
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                MISSING.append(str(path))
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Figure {self.fig_n}. {caption}")
        r.italic = True
        r.font.size = Pt(9)

    def shot(self, key, shots_dir, width_in=5.6):
        stem, caption = SHOTS[key]
        found = None
        for ext in ("png", "jpg", "jpeg", "PNG", "JPG"):
            cand = shots_dir / f"{stem}.{ext}"
            if cand.exists():
                found = cand
                break
        self.fig_n += 1
        if found:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(found), width=Inches(width_in))
        else:
            self._placeholder(f"Screenshot needed: screenshots/{stem}.png")
            MISSING.append(f"screenshots/{stem}.png")
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Figure {self.fig_n}. {caption}")
        r.italic = True
        r.font.size = Pt(9)

    def _placeholder(self, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        shade(c, "FFF9C4")
        set_cell_borders(c, "E0C200")
        para = c.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(14)
        para.paragraph_format.space_after = Pt(14)
        run = para.add_run(text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x80, 0x60, 0x00)

    def page_break(self):
        self.doc.add_page_break()


# ----------------------------------------------------------------------------------
# architecture diagram
# ----------------------------------------------------------------------------------
def make_architecture(out: Path, svc_port: str):
    fig, ax = plt.subplots(figsize=(11, 6.2))
    ax.set_xlim(0, 110)
    ax.set_ylim(0, 67)
    ax.axis("off")

    def box(x, y, w, h, text, fc="#EAF1FB", ec="#1F3A5F", fs=8.5, bold=False):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.2",
                                    fc=fc, ec=ec, lw=1.2))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs,
                fontweight="bold" if bold else "normal", color="#1A1A1A")

    def group(x, y, w, h, title, fc):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.3,rounding_size=1.5",
                                    fc=fc, ec="#8899AA", lw=1, ls="--"))
        ax.text(x + 1.2, y + h - 2.2, title, fontsize=9, fontweight="bold", color="#33475B")

    def arrow(x1, y1, x2, y2, label=None, ls="-"):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12,
                                     lw=1.3, color="#33475B", linestyle=ls))
        if label:
            ax.text((x1 + x2) / 2, (y1 + y2) / 2 + 1.3, label, fontsize=7.5, ha="center",
                    color="#33475B")

    # ---- build / CI zone
    group(1, 33, 52, 27, "Development and CI/CD", "#F7FAFD")
    box(3, 44, 13, 9, "Developer\nlaptop", fc="#FFFFFF")
    box(21, 44, 13, 9, "GitHub\nrepository", fc="#FFFFFF")
    box(39, 44, 12, 9, "GitHub\nActions", fc="#FFFFFF", bold=True)
    box(21, 35, 30, 6, "lint  ->  pytest  ->  train  ->  artifacts", fc="#FFF4DD", ec="#B07A00", fs=8)
    arrow(16.5, 48.5, 20.5, 48.5, "git push")
    arrow(34.5, 48.5, 38.5, 48.5, "trigger")
    arrow(45, 43.5, 45, 41.2)

    # ---- ML pipeline
    group(1, 2, 52, 28, "ML pipeline (local, reproducible)", "#F6FBF4")
    box(3, 15, 14, 9, "UCI Heart\nDisease data\n(download script)", fc="#FFFFFF", fs=8)
    box(21, 15, 14, 9, "Clean + EDA\n+ feature\npipeline", fc="#FFFFFF", fs=8)
    box(39, 15, 12, 9, "Train LR / RF\n5-fold CV\ntuning", fc="#FFFFFF", fs=8)
    box(21, 4, 14, 7, "MLflow tracking\n(params, metrics,\nplots)", fc="#FFFFFF", fs=7.5)
    box(39, 4, 12, 7, "model.joblib\n+ metadata", fc="#FFF4DD", ec="#B07A00", fs=8, bold=True)
    arrow(17.3, 19.5, 20.7, 19.5)
    arrow(35.3, 19.5, 38.7, 19.5)
    arrow(45, 14.5, 45, 11.2)
    arrow(40, 15, 35.3, 9.5, ls="--")

    # ---- runtime zone
    group(58, 2, 51, 58, "Local Kubernetes cluster (Rancher Desktop)", "#FBF7FB")
    box(60, 46, 15, 8, "Client\n(curl / Swagger)", fc="#FFFFFF")
    group(60, 22, 21, 20, "namespace: heart-api", "#FFFFFF")
    box(62, 34, 17, 5, f"LoadBalancer Svc :{svc_port}", fc="#EAF1FB", fs=7.5)
    box(62, 25, 7.5, 6, "Pod 1\nFastAPI", fc="#E4F5E4", ec="#2E7D32", fs=7.5)
    box(71, 25, 7.5, 6, "Pod 2\nFastAPI", fc="#E4F5E4", ec="#2E7D32", fs=7.5)
    box(62, 14, 17, 5, "Headless Svc (pod IPs)", fc="#EAF1FB", fs=7.5)
    group(84, 6, 23, 36, "namespace: monitoring", "#FFFFFF")
    box(86, 28, 19, 7, "Prometheus\n(scrape /metrics)", fc="#FDE8E4", ec="#C0392B", fs=8)
    box(86, 14, 19, 7, "Grafana\ndashboard", fc="#FFF1DC", ec="#E67E22", fs=8)
    arrow(67.5, 45.5, 70.5, 39.6, "POST /predict")
    arrow(70.5, 33.6, 65.7, 31.4)
    arrow(70.5, 33.6, 74.7, 31.4)
    arrow(70.5, 24.6, 70.5, 19.4, ls="--")
    arrow(79.5, 16.5, 85.5, 30.0, "scrape", ls="--")
    arrow(95.5, 27.5, 95.5, 21.4, "query")
    # model -> pod, CI -> image
    arrow(51.8, 7.5, 62, 24.5)
    ax.text(63, 8, "Docker image\n(contains the model)", fontsize=7.5, ha="left", color="#33475B")

    ax.text(55, 65, "Heart Disease Risk Prediction: MLOps architecture", ha="center",
            fontsize=11, fontweight="bold", color="#1F3A5F")
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ----------------------------------------------------------------------------------
# data gathering
# ----------------------------------------------------------------------------------
def load_facts(root: Path):
    f = {"clean": None, "cmp": None, "meta": None, "tests": []}
    clean = root / "data" / "processed" / "heart_clean.csv"
    if clean.exists():
        f["clean"] = pd.read_csv(clean)
    cmp_ = root / "reports" / "model_comparison.csv"
    if cmp_.exists():
        f["cmp"] = pd.read_csv(cmp_)
    meta = root / "models" / "model_metadata.json"
    if meta.exists():
        f["meta"] = json.loads(meta.read_text())
    for tf in sorted((root / "tests").glob("test_*.py")):
        n = len(re.findall(r"^def test_", tf.read_text(), flags=re.M))
        f["tests"].append((tf.name, n))
    return f


def pretty(name):
    return {"logistic_regression": "Logistic Regression",
            "random_forest": "Random Forest"}.get(name, name)


# ----------------------------------------------------------------------------------
# report body
# ----------------------------------------------------------------------------------
def build(root: Path, out: Path, args):
    facts = load_facts(root)
    fig_dir = root / "reports" / "figures"
    shots = root / "screenshots"
    arch = fig_dir / "architecture.png"
    make_architecture(arch, args.svc_port)

    R = Report()
    d = R.doc
    repo = args.repo_url or "[TODO: paste GitHub repository URL]"

    # ---- title block
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("End-to-End MLOps for Heart Disease Risk Prediction")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    t2 = d.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run("Machine Learning Operations (AIMLCZG523) - Assignment 01")
    r.font.size = Pt(12)
    sid = args.student_id or "[TODO: student ID]"
    R.table(["Field", "Detail"], [
        ["Student", args.name],
        ["Student ID", sid],
        ["Course", "MLOps (AIMLCZG523), BITS Pilani WILP"],
        ["Date", date.today().strftime("%d %B %Y")],
        ["Code repository", repo],
        ["Demo video", args.video_url or "[TODO: paste video link or file name]"],
    ], widths=[4, 12.4])

    # ---- 1 introduction
    R.h1("1. Introduction and Problem Statement")
    R.p("Heart disease is a leading cause of death, and early risk screening from routine "
        "clinical measurements can help prioritise patients for further tests. This project "
        "builds a binary classifier on the UCI Heart Disease (Cleveland) dataset that predicts "
        "whether a patient has heart disease, and delivers it the way a production ML system "
        "would be delivered: reproducible, tested, tracked, containerised, deployed and monitored.")
    R.p("**Objectives**")
    R.bullets([
        "Acquire, clean and explore the data with a scripted, repeatable process.",
        "Train and tune two classifiers (Logistic Regression and Random Forest) with cross-validation.",
        "Track every experiment in MLflow and package the best model with its preprocessing pipeline.",
        "Automate linting, testing and training with GitHub Actions.",
        "Serve the model through a FastAPI service in Docker and deploy it on Kubernetes.",
        "Monitor the running service with structured logs, Prometheus metrics and a Grafana dashboard.",
    ])

    # ---- 2 setup
    R.h1("2. Setup and Installation")
    R.p("Prerequisites: Python 3.11 (pinned in .tool-versions), Git, Docker, and a local Kubernetes "
        "cluster (Rancher Desktop was used). All commands run from the repository root.")
    R.code(f"""git clone {repo if not repo.startswith('[') else '<repository-url>'}
cd heart-disease-mlops
python -m venv .venv && source .venv/bin/activate
pip install -r requirements-dev.txt

python src/download_data.py      # fetch UCI data -> data/raw/heart_raw.csv
python src/preprocess.py         # clean data     -> data/processed/heart_clean.csv
python src/train.py              # tune, track in MLflow, save models/model.joblib
python src/package_model.py      # metadata + MLflow-format copy + round-trip check
python src/predict.py            # inference on data/sample_input.json
pytest && flake8 src tests       # tests and lint
mlflow ui --backend-store-uri ./mlruns --port 5001""")
    R.table(["File", "Purpose"], [
        ["requirements.txt", "Pinned packages for training, tracking and the pipeline"],
        ["requirements-dev.txt", "Adds pytest, flake8, Jupyter and httpx for development and CI"],
        ["requirements-api.txt", "Minimal pinned packages for the serving container"],
        [".tool-versions", "Pins the Python version used"],
    ], widths=[5, 11.4], caption="Dependency files")
    R.p("Reproducibility was verified by running the whole pipeline, from dataset download to "
        "prediction, in a brand-new virtual environment created only from requirements.txt.")
    R.shot("clean_env", shots, 5.2)

    # ---- 3 EDA
    R.h1("3. Data Acquisition and Exploratory Data Analysis")
    df = facts["clean"]
    R.h2("3.1 Dataset and cleaning")
    if df is not None:
        n1 = int((df["target"] == 1).sum())
        n0 = int((df["target"] == 0).sum())
        R.p(f"The Cleveland subset of the UCI Heart Disease dataset has {len(df)} patient records "
            f"and {df.shape[1] - 1} predictors plus a binary target. After binarising the original "
            f"0-4 diagnosis column (0 = no disease, 1-4 = disease), there are {n0} patients without "
            f"disease ({n0 / len(df):.1%}) and {n1} with disease ({n1 / len(df):.1%}).")
    else:
        R.p("[TODO: dataset size and class balance. data/processed/heart_clean.csv was not found]")
    R.p("The data is downloaded by src/download_data.py, which assigns the column names (the "
        "source file has no header) and treats '?' as missing.")
    R.table(["Issue", "Columns", "Action", "Reason"], [
        ["Missing values", "ca, thal", "Filled with the most frequent value",
         "Categorical codes, only a handful of rows affected"],
        ["Multi-class target", "num", "Binarised into target (0 / 1)",
         "The task is risk prediction (disease present or not)"],
        ["Duplicates", "all", "Dropped if present", "Avoid leakage between train and test"],
        ["Scaling / encoding", "numeric / categorical",
         "Done inside the model pipeline, not here",
         "Same transformation at training and serving time"],
    ], widths=[3.2, 3, 5, 5.2], caption="Cleaning decisions", font=8.5)

    R.h2("3.2 Exploratory analysis")
    R.figure_pair(fig_dir / "class_balance.png", fig_dir / "missing_values.png",
                  "Class balance (left) and missing values in the raw data (right)")
    R.figure(fig_dir / "numeric_histograms.png", "Numeric feature distributions split by target", 5.8)
    R.figure(fig_dir / "correlation_heatmap.png", "Correlation heatmap of all features", 4.6)
    R.figure(fig_dir / "categorical_vs_target.png", "Categorical features against the target", 5.8)
    if df is not None:
        corr = df.corr()["target"].drop("target").sort_values(key=abs, ascending=False).head(6)
        R.table(["Feature", "Correlation with target"],
                [[k, f"{v:+.3f}"] for k, v in corr.items()],
                widths=[5, 6], caption="Six strongest correlations with the target")
    R.p("**Key observations**")
    R.bullets([
        "[TODO: write 3-4 observations in your own words, e.g. class balance and what that means "
        "for the choice of metrics]",
        "[TODO: which numeric features separate the two classes best in the histograms/boxplots]",
        "[TODO: which categorical features (cp, thal, ca, exang) show the clearest pattern]",
        "[TODO: any outliers (e.g. chol) and whether you kept them, and why]",
    ])

    # ---- 4 modelling
    R.h1("4. Feature Engineering and Model Development")
    R.h2("4.1 Feature pipeline")
    R.p("All transformations live in one scikit-learn ColumnTransformer so the exact same steps "
        "run in training, tests and the API:")
    R.bullets([
        "**Numeric** (age, trestbps, chol, thalach, oldpeak): median imputation, then standard scaling.",
        "**Categorical** (sex, cp, fbs, restecg, exang, slope, ca, thal): most-frequent imputation, "
        "then one-hot encoding with handle_unknown='ignore'.",
        "Scaling matters for Logistic Regression because features such as chol and oldpeak sit on "
        "very different scales. Codes such as cp and thal are labels, not quantities, so they are "
        "one-hot encoded.",
    ])
    R.h2("4.2 Models, tuning and evaluation")
    R.p("The data was split 80/20 with stratification (random_state = 42). Each model was tuned "
        "with GridSearchCV using 5-fold stratified cross-validation on the training portion, "
        "optimising ROC-AUC. The winner was chosen on its cross-validation score, so the held-out "
        "test set never influenced model selection.")
    R.table(["Model", "Hyperparameter grid"], [[k, v] for k, v in GRIDS.items()],
            widths=[4, 12.4], caption="Search space")
    cmp_ = facts["cmp"]
    if cmp_ is not None:
        cols = ["model", "cv_roc_auc", "test_accuracy", "test_precision", "test_recall",
                "test_f1", "test_roc_auc"]
        rows = []
        for _, row in cmp_.iterrows():
            rows.append([pretty(row["model"])] + [f"{row[c]:.3f}" for c in cols[1:]])
        R.table(["Model", "CV ROC-AUC", "Accuracy", "Precision", "Recall", "F1", "Test ROC-AUC"],
                rows, widths=[3.6, 2.2, 2.0, 2.2, 2.0, 1.8, 2.6],
                caption="Cross-validation and held-out test results")
        best = cmp_.iloc[0]
        R.p(f"**Selected model: {pretty(best['model'])}** with best parameters "
            f"{best['best_params']}, CV ROC-AUC {best['cv_roc_auc']:.3f} and test recall "
            f"{best['test_recall']:.3f}.")
    else:
        R.p("[TODO: reports/model_comparison.csv not found. Run python src/train.py]")
    R.figure_pair(fig_dir / "roc_logistic_regression.png", fig_dir / "roc_random_forest.png",
                  "ROC curves on the test set: Logistic Regression (left), Random Forest (right)")
    R.figure_pair(fig_dir / "cm_logistic_regression.png", fig_dir / "cm_random_forest.png",
                  "Confusion matrices on the test set")
    R.p("[TODO: 3-4 sentences on why you picked this model: compare CV score vs test score, "
        "recall (missed disease cases are costly), interpretability, and any overfitting you saw]")

    # ---- 5 MLflow
    R.h1("5. Experiment Tracking")
    R.p("MLflow records every training run so results can be compared, audited and reproduced. "
        "Each model gets a parent run, and every hyperparameter combination tried by the grid "
        "search is logged as a nested child run.")
    R.table(["Logged item", "Details"], [
        ["Parameters", "model name, split size, CV folds, random state, row counts, search grid, best parameters"],
        ["Metrics", "per-candidate CV ROC-AUC mean and std; best CV ROC-AUC; test accuracy, precision, recall, F1, ROC-AUC"],
        ["Artifacts", "confusion matrix and ROC curve (plots/), full pipeline (model/), model_comparison.csv"],
        ["Tags", "model_type, stage, selected_as_best on the winning run"],
    ], widths=[3.2, 13.2], caption="What is logged to MLflow", font=8.5)
    R.shot("mlflow_runs", shots)
    R.shot("mlflow_compare", shots)
    R.shot("mlflow_artifacts", shots, 4.8)
    R.p("[TODO: 2-3 sentences on what the candidate runs showed, e.g. how much C or max_depth "
        "changed the CV ROC-AUC]")

    # ---- 6 packaging
    R.h1("6. Model Packaging and Reproducibility")
    R.p("The final artefact is a single scikit-learn Pipeline (imputers, scaler, one-hot encoder "
        "and classifier) saved as models/model.joblib, plus an MLflow-format copy in "
        "models/mlflow_model that carries its own environment files. Because preprocessing is part "
        "of the saved object, the API receives raw feature values and cannot drift from the "
        "training-time transformations. joblib was chosen over ONNX because it keeps the complete "
        "pipeline, including one-hot encoding, without conversion tooling.")
    meta = facts["meta"]
    if meta:
        R.table(["Metadata field", "Value"], [
            ["Model", pretty(meta.get("model_name", ""))],
            ["Trained at (UTC)", meta.get("trained_at_utc", "")],
            ["Python / scikit-learn", f"{meta.get('python_version', '')} / {meta.get('scikit_learn_version', '')}"],
            ["Training data hash (SHA-256, truncated)", meta.get("training_data_sha256", "")],
            ["Model file hash (SHA-256, truncated)", meta.get("model_sha256", "")],
            ["Random state", meta.get("random_state", "")],
        ], widths=[7, 9.4], caption="Model metadata (models/model_metadata.json)", font=8.5)
    else:
        R.p("[TODO: models/model_metadata.json not found. Run python src/package_model.py]")
    R.p("**Reproducibility measures**")
    R.bullets([
        "Fixed random_state = 42 for the split, cross-validation and models; training twice "
        "produced an identical comparison table.",
        "Pinned requirements files and a pinned Python version.",
        "package_model.py checks that the joblib and MLflow copies give identical predictions.",
        "The pipeline was run end-to-end in a clean virtual environment (Figure 1).",
    ])
    R.shot("predict_cli", shots, 4.6)

    # ---- 7 CI/CD
    R.h1("7. CI/CD Pipeline and Automated Testing")
    if facts["tests"]:
        desc = {
            "test_preprocess.py": "Target binarisation, missing-value handling, duplicates, no input mutation",
            "test_features.py": "Scaling, one-hot encoding, unseen categories, imputation",
            "test_model.py": "Metric ranges, probability validity, training reproducibility",
            "test_predict.py": "Output format, confidence range, missing-feature error, saved model",
            "test_api.py": "Health, /predict, validation errors (422), /metrics endpoint",
        }
        rows = [[n, desc.get(n, ""), c] for n, c in facts["tests"]]
        rows.append(["Total", "", sum(c for _, c in facts["tests"])])
        R.table(["Test file", "What it covers", "Tests"], rows, widths=[3.6, 10.8, 2],
                caption="Unit-test suite (pytest)", font=8.5)
    R.p("The tests use small synthetic datasets so they run in seconds and do not depend on "
        "local files. The GitHub Actions workflow (.github/workflows/ci.yml) has three chained jobs:")
    R.table(["Job", "Steps", "Artifacts"], [
        ["lint", "flake8 on src/ and tests/", "flake8-report.txt"],
        ["test", "pytest (needs lint)", "pytest-report.xml"],
        ["train", "preprocess, train with MLflow, package model, smoke-test inference (needs test)",
         "models/, reports/, mlruns/"],
    ], widths=[2.5, 9.5, 4.4], caption="CI/CD workflow", font=8.5)
    R.p("A failure in lint or test stops the pipeline and skips later jobs, so broken code never "
        "reaches the training stage. Logs for every step are retained by GitHub.")
    R.shot("local_pytest", shots, 4.8)
    R.shot("ci_green", shots)
    R.shot("ci_pytest", shots, 4.8)
    R.shot("ci_artifacts", shots, 4.8)

    # ---- 8 containerisation & deployment
    R.h1("8. Containerisation and Deployment")
    R.h2("8.1 Serving API")
    R.p("The service is built with FastAPI. POST /predict accepts the 13 patient features as JSON "
        "and returns the prediction and a confidence value. Pydantic validates every field "
        "(type and range), so invalid input is rejected with HTTP 422 before it reaches the model. "
        "GET /health supports Kubernetes probes and GET /metrics exposes Prometheus metrics.")
    R.code("""Request  POST /predict
{"age": 63, "sex": 1, "cp": 1, "trestbps": 145, "chol": 233, "fbs": 1, "restecg": 2,
 "thalach": 150, "exang": 0, "oldpeak": 2.3, "slope": 3, "ca": 0, "thal": 6}

Response 200
{"prediction": 1, "label": "disease", "confidence": 0.87, "probability_disease": 0.87}""")
    R.p("[TODO: replace the response above with the real output from your container]")
    R.shot("swagger_local", shots, 5.0)
    R.h2("8.2 Docker image")
    R.bullets([
        "Slim Python 3.11 base image; dependencies installed before the code so Docker caches that layer.",
        "A separate requirements-api.txt keeps MLflow, Matplotlib and Jupyter out of the image.",
        "Runs as a non-root user and includes a HEALTHCHECK.",
        ".dockerignore keeps the virtual environment, data, notebooks and tests out of the build context.",
    ])
    R.h2("8.3 Kubernetes deployment")
    R.p("The image was deployed on a local Kubernetes cluster (Rancher Desktop) using manifests "
        "in k8s/ and exposed through a LoadBalancer service:")
    R.table(["Manifest", "Defines"], [
        ["00-namespace.yaml", "heart-api namespace"],
        ["10-deployment.yaml", "2 replicas, rolling update (maxUnavailable 0), readiness and liveness probes on /health, "
                               "CPU/memory requests and limits, non-root security context"],
        ["20-service.yaml", f"LoadBalancer service on port {args.svc_port} forwarding to container port 8000"],
        ["30-service-headless.yaml", "Headless service so Prometheus scrapes each pod separately"],
    ], widths=[4.6, 11.8], caption="Kubernetes manifests", font=8.5)
    R.p(f"Access: kubectl apply -f k8s/, then open http://localhost:{args.svc_port}/docs. "
        "The following were demonstrated: traffic shared across both pods, automatic replacement "
        "of a deleted pod, scaling, and a rolling update with rollback.")
    R.shot("k8s_get_all", shots, 5.2)
    R.shot("k8s_describe", shots, 5.0)

    # ---- 9 monitoring
    R.h1("9. Monitoring and Logging")
    R.p("Every request passes through a middleware that writes one structured JSON log line "
        "(request ID, method, path, status, duration) and updates Prometheus metrics. Patient "
        "input values are deliberately not logged, to avoid writing health data to logs. "
        "Health and metrics calls are excluded so probes do not flood the logs.")
    R.code('{"event": "request", "request_id": "3f9a1c2e", "method": "POST", '
           '"path": "/predict", "status": 200, "duration_ms": 6.41}')
    R.table(["Metric", "Type", "Purpose"], [
        ["api_requests_total{method,path,status}", "Counter", "Traffic and error rates"],
        ["api_request_duration_seconds{path}", "Histogram", "Latency percentiles (p95)"],
        ["api_predictions_total{label}", "Counter", "Predicted class mix, a simple drift signal"],
        ["api_prediction_confidence", "Histogram", "Model confidence over time"],
    ], widths=[6.6, 2.4, 7.4], caption="Exported metrics", font=8.5)
    R.p("Prometheus discovers both API pods through the headless service and scrapes them every "
        "10 seconds. The Grafana dashboard shows request rate, p95 latency, pods up, 5xx error "
        "rate, average confidence and predictions by class. It was tested with 300 generated "
        "requests, about 10% of them deliberately invalid.")
    R.shot("grafana", shots)
    R.shot("k8s_logs", shots, 5.0)
    R.shot("metrics_raw", shots, 4.8)
    R.shot("pods_all", shots, 4.8)

    # ---- 10 architecture & conclusion
    R.h1("10. Architecture, Repository and Conclusion")
    R.figure(arch, "End-to-end architecture", 6.3)
    R.h2("10.1 Repository structure")
    R.code("""heart-disease-mlops/
  .github/workflows/ci.yml     CI/CD pipeline
  data/                        raw + cleaned dataset, sample input
  notebooks/01_eda.ipynb       exploratory analysis
  src/                         download, preprocess, features, train, package, predict, app
  tests/                       pytest unit tests
  models/                      model.joblib, metadata, MLflow-format copy
  k8s/  monitoring/            Kubernetes and Prometheus/Grafana manifests
  scripts/                     traffic generator, report builder
  screenshots/  reports/       evidence and figures
  Dockerfile  requirements*.txt""")
    R.p(f"**Repository:** {repo}")
    R.p(f"**Access:** the API runs locally at http://localhost:{args.svc_port}/docs after "
        "kubectl apply -f k8s/. No public URL is exposed.")
    R.h2("10.2 Conclusion and limitations")
    R.p("The project delivers a complete, automated path from raw data to a monitored, "
        "containerised prediction service.")
    R.bullets([
        "The dataset is small (303 rows), so metrics have wide uncertainty and the model is a "
        "learning exercise, not a clinical tool.",
        "The model is a static artefact; there is no automated retraining or drift alerting yet.",
        "Deployment is local; moving to a cloud cluster would need a container registry and Ingress."
    ])

    out.parent.mkdir(parents=True, exist_ok=True)
    R.doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=str(Path(__file__).resolve().parents[1]))
    ap.add_argument("--out", default=None)
    ap.add_argument("--name", default="Abdullah Farhan Safwi")
    ap.add_argument("--student-id", default="")
    ap.add_argument("--repo-url", default="")
    ap.add_argument("--video-url", default="")
    ap.add_argument("--svc-port", default="8081")
    args = ap.parse_args()

    root = Path(args.root)
    out = Path(args.out) if args.out else root / "reports" / "MLOps_Assignment1_Report.docx"
    build(root, out, args)

    print(f"\nReport written: {out}")
    print(f"Yellow [TODO] items to fill in: {TODO_COUNT}")
    if MISSING:
        print(f"\nMissing images ({len(MISSING)}):")
        for m in dict.fromkeys(MISSING):
            print("  -", m)
    else:
        print("All images found.")


if __name__ == "__main__":
    main()